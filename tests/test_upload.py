from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas

from app.db.session import SessionLocal
from app.main import app
from app.models.document import DocumentChunk


client = TestClient(app)


def create_test_docx() -> bytes:
    document = Document()

    document.add_paragraph(
        "Employees receive twenty days of annual leave."
    )

    document.add_paragraph(
        "Employees can request additional leave according to company policy."
    )

    output = BytesIO()
    document.save(output)

    return output.getvalue()


def create_test_pdf() -> bytes:
    writer = PdfWriter()

    writer.add_blank_page(
        width=612,
        height=792,
    )

    output = BytesIO()
    writer.write(output)

    return output.getvalue()


def create_multi_page_test_pdf() -> bytes:
    output = BytesIO()

    pdf = canvas.Canvas(output)

    pdf.drawString(
        72,
        750,
        "Employees receive twenty days of annual leave.",
    )

    pdf.showPage()

    pdf.drawString(
        72,
        750,
        "Employees can request additional leave according to company policy.",
    )

    pdf.save()

    return output.getvalue()


def test_upload_txt_file():
    response = client.post(
        "/documents/upload",
        files={
            "file": (
                "employee_policy.txt",
                (
                    "Employees receive twenty days of annual leave."
                ).encode("utf-8"),
                "text/plain",
            )
        },
    )

    assert response.status_code == 200

    data = response.json()

    assert data["filename"] == "employee_policy.txt"
    assert data["content_type"] == "text/plain"
    assert data["file_size"] > 0
    assert data["status"] == "uploaded"


def test_upload_docx_file():
    content = create_test_docx()

    response = client.post(
        "/documents/upload",
        files={
            "file": (
                "employee_policy.docx",
                content,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 200

    data = response.json()

    assert data["filename"] == "employee_policy.docx"
    assert data["file_size"] == len(content)
    assert data["status"] == "uploaded"


def test_upload_pdf_file():
    content = create_test_pdf()

    response = client.post(
        "/documents/upload",
        files={
            "file": (
                "employee_policy.pdf",
                content,
                "application/pdf",
            )
        },
    )

    assert response.status_code == 400

    assert (
        response.json()["detail"]
        == "No text could be extracted from the uploaded document"
    )


def test_upload_multi_page_pdf_preserves_page_metadata():
    content = create_multi_page_test_pdf()

    response = client.post(
        "/documents/upload",
        files={
            "file": (
                "employee_policy_pages.pdf",
                content,
                "application/pdf",
            )
        },
    )

    assert response.status_code == 200

    data = response.json()

    assert data["filename"] == "employee_policy_pages.pdf"
    assert data["status"] == "uploaded"

    document_id = data["id"]

    db = SessionLocal()

    try:
        chunks = (
            db.query(DocumentChunk)
            .filter(
                DocumentChunk.document_id == document_id
            )
            .order_by(DocumentChunk.chunk_index)
            .all()
        )

        assert len(chunks) == 2

        first_chunk = chunks[0]
        second_chunk = chunks[1]

        assert first_chunk.chunk_index == 0
        assert second_chunk.chunk_index == 1

        assert first_chunk.metadata_ == {
            "source": "employee_policy_pages.pdf",
            "page": 1,
            "chunk_index": 0,
        }

        assert second_chunk.metadata_ == {
            "source": "employee_policy_pages.pdf",
            "page": 2,
            "chunk_index": 1,
        }

        assert (
            "Employees receive twenty days"
            in first_chunk.content
        )

        assert (
            "Employees can request additional leave"
            in second_chunk.content
        )

    finally:
        db.close()


def test_upload_unsupported_file_type():
    response = client.post(
        "/documents/upload",
        files={
            "file": (
                "employee_policy.xlsx",
                b"fake excel content",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert response.status_code == 400

    assert (
        response.json()["detail"]
        == (
            "Unsupported file type. "
            "Supported formats: .txt, .pdf, .docx"
        )
    )


def test_upload_empty_file():
    response = client.post(
        "/documents/upload",
        files={
            "file": (
                "empty.txt",
                b"",
                "text/plain",
            )
        },
    )

    assert response.status_code == 400

    assert response.json()["detail"] == "Uploaded file is empty"