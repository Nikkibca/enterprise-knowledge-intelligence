from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfWriter
from reportlab.pdfgen import canvas

from app.services.document_extractor import (
    DOCXDocumentExtractor,
    PDFDocumentExtractor,
    TextDocumentExtractor,
)


def test_text_document_extractor_extracts_utf8_text():
    extractor = TextDocumentExtractor()

    content = "Employees receive twenty days of annual leave.".encode(
        "utf-8"
    )

    result = extractor.extract(content)

    assert result == (
        "Employees receive twenty days of annual leave."
    )


def test_text_document_extractor_supports_unicode():
    extractor = TextDocumentExtractor()

    content = "Employee policy — annual leave".encode("utf-8")

    result = extractor.extract(content)

    assert result == "Employee policy — annual leave"


def test_text_document_extractor_rejects_invalid_utf8():
    extractor = TextDocumentExtractor()

    invalid_content = b"\xff\xfe\xfa"

    with pytest.raises(UnicodeDecodeError):
        extractor.extract(invalid_content)


def create_test_pdf() -> bytes:
    output = BytesIO()

    pdf = canvas.Canvas(output)

    pdf.drawString(
        72,
        750,
        "Employees receive twenty days of annual leave.",
    )

    pdf.drawString(
        72,
        730,
        "Employees can request additional leave according to company policy.",
    )

    pdf.save()

    return output.getvalue()


def test_pdf_document_extractor_reads_pdf_text():
    extractor = PDFDocumentExtractor()

    content = create_test_pdf()

    result = extractor.extract(content)

    assert "Employees receive twenty days of annual leave." in result

    assert (
        "Employees can request additional leave according to company policy."
        in result
    )

def test_pdf_document_extractor_preserves_page_metadata():
    extractor = PDFDocumentExtractor()

    content = create_test_pdf()

    blocks = extractor.extract_structured(content)

    assert len(blocks) == 1

    assert blocks[0].metadata == {
        "page": 1,
    }

    assert (
        "Employees receive twenty days of annual leave."
        in blocks[0].text
    )


def create_blank_test_pdf() -> bytes:
    writer = PdfWriter()

    writer.add_blank_page(
        width=612,
        height=792,
    )

    output = BytesIO()
    writer.write(output)

    return output.getvalue()


def test_pdf_document_extractor_handles_blank_pdf():
    extractor = PDFDocumentExtractor()

    content = create_blank_test_pdf()

    result = extractor.extract(content)

    assert result == ""


def create_test_docx() -> bytes:
    document = Document()

    document.add_paragraph(
        "Employees receive twenty days of annual leave."
    )

    document.add_paragraph(
        "Employees can request additional leave according to company policy."
    )

    output = BytesIO()
    document.save(output)

    return output.getvalue()


def test_docx_document_extractor_reads_docx():
    extractor = DOCXDocumentExtractor()

    content = create_test_docx()

    result = extractor.extract(content)

    assert "Employees receive twenty days of annual leave." in result

    assert (
        "Employees can request additional leave according to company policy."
        in result
    )